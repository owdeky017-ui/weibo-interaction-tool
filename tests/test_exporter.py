"""``exporter`` 测试：Excel 三表结构、CSV BOM、长文本不截断、空记录分支。

导出去掉 pandas 之后，Excel 的分块排序、汇总表、列宽/冻结这些细节全部
由自己维护，属于「重构风险点」，因此逐项固定下来。
"""

from __future__ import annotations

import csv
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

import exporter

LONG_TEXT = "长" * 300  # 300 字：验证不再被截断到 120 字

USER_A = {"screen_name": "甲", "uid": "1"}
USER_B = {"screen_name": "乙", "uid": "2"}

Rec = Callable[..., dict[str, Any]]


@pytest.fixture
def records(rec: Rec) -> list[dict[str, Any]]:
    """6 条记录：A→B 3 条（转发/点赞/@提及）、B→A 3 条（评论/评论回复/转发）。"""
    return [
        rec("转发", "A→B", "甲", LONG_TEXT, "转发语"),
        rec("点赞", "A→B", "甲", "短内容", "点赞了这条微博"),
        rec("评论", "B→A", "乙", "另一条", "评论内容"),
        rec("评论回复", "B→A", "乙", "第三条", "回复内容"),
        rec("转发", "B→A", "乙", "第四条", "转发语2"),
        rec("@提及", "A→B", "甲", "第五条", "@乙"),
    ]


@pytest.fixture
def exported(records: list[dict[str, Any]], tmp_path: Path) -> dict[str, Any]:
    return exporter.export(records, USER_A, USER_B, out_dir=str(tmp_path), formats={"html", "excel", "csv"})


# --------------------------------------------------------------------------- #
# 返回结构与文件落盘
# --------------------------------------------------------------------------- #


class TestExportResult:
    def test_count(self, exported: dict[str, Any]) -> None:
        assert exported["count"] == 6

    @pytest.mark.parametrize("fmt", ["excel", "csv", "html"])
    def test_file_created(self, exported: dict[str, Any], fmt: str) -> None:
        assert exported.get(fmt), f"未返回 {fmt} 路径"
        assert Path(exported[fmt]).exists(), f"{fmt} 文件不存在"

    def test_only_requested_formats_are_produced(self, records: list[dict[str, Any]], tmp_path: Path) -> None:
        res = exporter.export(records, USER_A, USER_B, out_dir=str(tmp_path), formats={"csv"})
        assert "csv" in res
        assert "excel" not in res
        assert "html" not in res

    def test_html_is_self_contained(self, exported: dict[str, Any]) -> None:
        """导出的 HTML 必须能离线打开：不能引用外部 http(s) 资源。"""
        content = Path(exported["html"]).read_text(encoding="utf-8")
        assert 'src="http' not in content
        assert '<link rel="stylesheet" href="http' not in content


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #


class TestExcelStructure:
    def test_sheet_order(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        assert wb.sheetnames == ["A→B", "B→A", "汇总统计"]

    def test_header_matches_columns(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        assert [c.value for c in wb["A→B"][1]] == exporter.COLUMNS

    def test_row_counts_per_direction(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        assert wb["A→B"].max_row == 4, "3 条数据 + 1 行表头"
        assert wb["B→A"].max_row == 4

    def test_first_row_is_frozen(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        assert wb["A→B"].freeze_panes == "A2"

    def test_column_width_is_set(self, exported: dict[str, Any]) -> None:
        """微博内容列要够宽，否则打开就是一堆 ####。"""
        wb = load_workbook(exported["excel"])
        assert (wb["A→B"].column_dimensions["F"].width or 0) > 10

    def test_long_text_not_truncated(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        assert wb["A→B"].cell(row=2, column=6).value == LONG_TEXT

    def test_all_columns_written(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        row = [c.value for c in wb["A→B"][2]]
        assert len(row) == len(exporter.COLUMNS)
        # 前 6 列（时间 / 互动类型 / 方向 / 发起方 / 微博作者 / 微博内容）必有值。
        # 后 6 列在部分记录里本就是空串，openpyxl 读回来是 None，属正常。
        assert all(row[:6]), f"关键列出现空值：{row}"


class TestPivotSheet:
    def test_header(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        assert [c.value for c in wb["汇总统计"][1]] == ["互动类型", "A→B", "B→A", "合计"]

    def test_last_row_is_total(self, exported: dict[str, Any]) -> None:
        wb = load_workbook(exported["excel"])
        rows = [[c.value for c in r] for r in wb["汇总统计"].iter_rows(min_row=2)]
        total = rows[-1]
        assert total[0] == "总计"
        assert total[1] == 3
        assert total[2] == 3
        assert total[3] == 6

    def test_returned_pivot_matches_sheet(self, exported: dict[str, Any]) -> None:
        assert exported["pivot"]["columns"] == ["互动类型", "A→B", "B→A", "合计"]
        assert exported["pivot"]["rows"][-1][0] == "总计"


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #


class TestCsv:
    def test_header(self, exported: dict[str, Any]) -> None:
        with open(exported["csv"], encoding="utf-8-sig", newline="") as f:
            assert next(csv.reader(f)) == exporter.COLUMNS

    def test_row_count(self, exported: dict[str, Any]) -> None:
        with open(exported["csv"], encoding="utf-8-sig", newline="") as f:
            assert len(list(csv.reader(f))) == 7, "1 行表头 + 6 行数据"

    def test_has_utf8_bom(self, exported: dict[str, Any]) -> None:
        """带 BOM，Excel 直接双击打开中文不乱码。"""
        with open(exported["csv"], "rb") as f:
            assert f.read(3) == b"\xef\xbb\xbf"

    def test_long_text_not_truncated(self, exported: dict[str, Any]) -> None:
        with open(exported["csv"], encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        assert LONG_TEXT in [r[5] for r in rows[1:]]


# --------------------------------------------------------------------------- #
# 空记录分支
# --------------------------------------------------------------------------- #


class TestEmptyRecords:
    @pytest.fixture
    def empty(self, tmp_path: Path) -> dict[str, Any]:
        return exporter.export([], USER_A, USER_B, out_dir=str(tmp_path), formats={"html", "excel", "csv"})

    def test_count_is_zero(self, empty: dict[str, Any]) -> None:
        assert empty["count"] == 0

    def test_excel_single_sheet_named_interaction_records(self, empty: dict[str, Any]) -> None:
        wb = load_workbook(empty["excel"])
        assert wb.sheetnames == ["互动记录"]

    def test_excel_header_still_written(self, empty: dict[str, Any]) -> None:
        wb = load_workbook(empty["excel"])
        assert [c.value for c in wb["互动记录"][1]] == exporter.COLUMNS

    def test_csv_has_header_only(self, empty: dict[str, Any]) -> None:
        with open(empty["csv"], encoding="utf-8-sig", newline="") as f:
            assert len(list(csv.reader(f))) == 1

    def test_html_still_generated(self, empty: dict[str, Any]) -> None:
        assert Path(empty["html"]).exists()

    def test_pivot_is_none(self, empty: dict[str, Any]) -> None:
        assert empty["pivot"] is None


# --------------------------------------------------------------------------- #
# 依赖收敛
# --------------------------------------------------------------------------- #


class TestNoPandas:
    def test_source_has_no_pandas_usage(self) -> None:
        """去掉 pandas 后不应再残留任何 ``pd.`` 调用（否则打包体积白省）。"""
        src = (Path(exporter.__file__)).read_text(encoding="utf-8")
        assert "import pandas" not in src
        assert "pd." not in src

    def test_source_does_not_import_analyzer(self) -> None:
        """依赖方向必须单向：exporter 不该反向 import analyzer。

        领域类型统一放在 models.py，这是重构时特意立下的约束。
        """
        src = (Path(exporter.__file__)).read_text(encoding="utf-8")
        assert "from analyzer import" not in src
        assert "import analyzer" not in src
