"""冻结环境自检：在**打包后的运行时**里验证关键依赖真的能用。

为什么需要它：PyInstaller 把模块打进 exe 只能证明「文件在」，
不能证明「运行时可用」——证书链缺了 cacert.pem、sqlite3.dll 版本不对、
PIL.ImageTk 找不到 _imaging 之类的问题，只有真跑一次才暴露。

用与 gui_app 完全相同的 PyInstaller 选项编译（唯一区别是控制台模式，
好让结果能打印出来），然后运行。

用法（由 build.py --probe 自动调用，也可单独跑）：
    python build.py --probe
"""

import os
import sys
import tempfile

# 冻结环境下 stdout 默认走系统代码页（中文 Windows 是 GBK），
# 调用方按 UTF-8 读取会变乱码，这里统一强制成 UTF-8。
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

RESULTS = []


def chk(name, fn):
    try:
        extra = fn()
        RESULTS.append((True, name, extra or ""))
    except Exception as e:
        import traceback

        RESULTS.append((False, name, f"{type(e).__name__}: {e}"))
        print(traceback.format_exc(), file=sys.stderr)


# ---------- 1. 全部依赖模块导入 ----------
def _imports():
    import sqlite3
    import tkinter

    import openpyxl
    import requests
    from PIL import Image

    import build_mode

    return (
        f"requests {requests.__version__} / openpyxl {openpyxl.__version__} / "
        f"sqlite3 {sqlite3.sqlite_version} / tk {tkinter.TkVersion} / "
        f"PIL {Image.__version__} / A_MODE={build_mode.A_MODE!r}"
    )


chk("导入全部依赖模块", _imports)


# ---------- 2. HTTPS + CA 证书链 ----------
def _https():
    import requests

    r = requests.get("https://weibo.com/", timeout=20, allow_redirects=False)
    assert r.status_code < 500, f"HTTP {r.status_code}"
    return f"HTTP {r.status_code}，{len(r.content)} 字节（SSL 校验通过）"


chk("HTTPS 请求 + CA 证书链", _https)


# ---------- 3. SQLite 断点读写 ----------
def _sqlite():
    from checkpoint import CheckpointStore

    d = tempfile.mkdtemp(prefix="probe_cp_")
    st = CheckpointStore(os.path.join(d, "cp.db"))
    rec = {"时间": "t", "互动类型": "点赞", "方向": "A→B", "_key": ["k"]}
    st.save([rec], {"m1"}, 123.0, {"x": 1}, True, [])
    got = st.load()
    assert len(got["records"]) == 1, got
    assert got["scanned_mids"] == ["m1"], got
    assert got["like_api_ok"] is True, got
    return f"写入并读回 {len(got['records'])} 条记录 / 1 个 mid"


chk("SQLite 断点读写", _sqlite)


# ---------- 4. Excel / CSV / HTML 导出 ----------
def _excel():
    from openpyxl import load_workbook

    from exporter import export

    d = tempfile.mkdtemp(prefix="probe_xl_")
    rec = {
        "时间": "2026-01-01 00:00:00",
        "互动类型": "转发",
        "方向": "A→B",
        "发起方": "甲",
        "微博作者": "乙",
        "微博内容": "内容",
        "互动内容": "转发语",
        "被回复人": "",
        "被回复评论": "",
        "微博时间": "2026-01-01 00:00:00",
        "微博链接": "u",
        "互动链接": "u",
        "_key": ["k"],
    }
    res = export(
        [rec],
        {"screen_name": "甲", "uid": "1"},
        {"screen_name": "乙", "uid": "2"},
        out_dir=d,
        formats={"html", "excel", "csv"},
    )
    wb = load_workbook(res["excel"])
    assert res["count"] == 1, res
    for k in ("excel", "csv", "html"):
        assert res.get(k) and os.path.exists(res[k]), k
    return f"xlsx sheet={wb.sheetnames}，count={res['count']}"


chk("Excel / CSV / HTML 导出", _excel)


# ---------- 5. tkinter + PIL.ImageTk（二维码显示路径） ----------
def _tk():
    import tkinter as tk

    from PIL import Image, ImageTk

    root = tk.Tk()
    root.withdraw()
    img = Image.new("RGB", (40, 40), "white")
    photo = ImageTk.PhotoImage(img)
    tk.Label(root, image=photo)
    root.update_idletasks()
    root.destroy()
    return "Tk 窗口 + PhotoImage 正常"


chk("tkinter + PIL.ImageTk", _tk)


# ---------- 6. frozen 路径解析 ----------
def _paths():
    import config

    assert getattr(sys, "frozen", False), "未运行在冻结环境（这是源码直跑？）"
    exe_dir = os.path.dirname(sys.executable)
    assert exe_dir == config.BASE_DIR, f"BASE_DIR={config.BASE_DIR} != exe 目录 {exe_dir}"
    assert os.path.join(exe_dir, "data") == config.DATA_DIR, config.DATA_DIR
    return f"BASE_DIR = {config.BASE_DIR}"


chk("frozen 路径解析（BASE_DIR = exe 目录）", _paths)


# ---------- 输出 ----------
print("=" * 62)
print("冻结环境自检 probe_frozen")
print("=" * 62)
print("Python :", sys.version.split()[0])
print("frozen :", getattr(sys, "frozen", False))
print("exe    :", sys.executable)
print("-" * 62)
for good, name, extra in RESULTS:
    print(f"[{'OK  ' if good else '失败'}] {name}" + (f"\n        —— {extra}" if extra else ""))
bad = [n for g, n, _ in RESULTS if not g]
print("-" * 62)
print(f"{len(RESULTS) - len(bad)} 通过 / {len(bad)} 失败")
for n in bad:
    print("  -", n)
sys.exit(1 if bad else 0)
